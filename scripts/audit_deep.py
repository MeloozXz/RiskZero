import docx

doc = docx.Document('TCC_Documentacao_Completa.docx')
print(f"Total de parágrafos: {len(doc.paragraphs)}")

for i, p in enumerate(doc.paragraphs):
    xml = p._element.xml
    if 'w:br' in xml and 'type="page"' in xml:
        prev_p = doc.paragraphs[i-1].text.strip() if i > 0 else 'START'
        curr_p = p.text.strip()
        next_p = doc.paragraphs[i+1].text.strip() if i+1 < len(doc.paragraphs) else 'END'
        print(f"PB at p[{i:03d}]: prev='{prev_p[:35]}' | curr='{curr_p[:35]}' | next='{next_p[:35]}'")
