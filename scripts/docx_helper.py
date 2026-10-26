import os
import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def get_base_doc_path():
    candidates = [
        os.path.join(os.path.dirname(__file__), "..", "TCC CERTO.docx"),
        os.path.join(os.path.dirname(__file__), "..", "seguran-aapp-main", "seguran-aapp-main", "TCC CERTO.docx"),
        os.path.join(os.getcwd(), "TCC CERTO.docx"),
        os.path.join(os.getcwd(), "seguran-aapp-main", "seguran-aapp-main", "TCC CERTO.docx"),
    ]
    for c in candidates:
        if os.path.exists(c):
            return os.path.abspath(c)
    raise FileNotFoundError("Não foi possível encontrar TCC CERTO.docx")

def get_target_doc_path():
    # Retorna o caminho para TCC_Documentacao_Completa.docx
    return os.path.join(os.path.dirname(__file__), "..", "TCC_Documentacao_Completa.docx")

def get_workspace_target_doc_path():
    return os.path.join("C:\\Users\\Aluno\\Downloads\\seguran-aapp-main (2)", "TCC_Documentacao_Completa.docx")

def set_cell_background(cell, hex_color):
    tcPr = cell._element.get_or_add_tcPr()
    # Remove existing shd if any
    for child in list(tcPr):
        if child.tag.endswith('shd'):
            tcPr.remove(child)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag.endswith('tcMar'):
            tcPr.remove(child)
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="D1D5DB", sz="4", val="single"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        for child in list(tblPr[0]):
            if child.tag.endswith('tblBorders'):
                tblPr[0].remove(child)
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:left w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideV w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def cant_split_row(row):
    trPr = row._element.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

def repeat_header_row(row):
    trPr = row._element.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.bold = True
    run.font.name = 'Arial'
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A) # Slate 900
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.bold = True
    run.font.name = 'Arial'
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B) # Slate 800
    return p

def add_heading_3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.bold = True
    run.font.name = 'Arial'
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x33, 0x41, 0x55) # Slate 700
    return p

def add_body_p(doc, text, bold_prefix=None, first_line_indent=1.25):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.space_before = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if first_line_indent:
        p.paragraph_format.first_line_indent = Cm(first_line_indent)
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.name = 'Arial'
        r_pre.font.size = Pt(12)
        r_pre.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
        
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
    return p

def add_bullet_p(doc, text, bold_title=None):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.left_indent = Cm(1.25)
    p.paragraph_format.first_line_indent = Cm(-0.5)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    # Marcador elegante
    r_bullet = p.add_run("• ")
    r_bullet.bold = True
    r_bullet.font.name = 'Arial'
    r_bullet.font.size = Pt(11)
    r_bullet.font.color.rgb = RGBColor(0x25, 0x63, 0xEB) # Azul Royal
    
    if bold_title:
        r_title = p.add_run(bold_title + " ")
        r_title.bold = True
        r_title.font.name = 'Arial'
        r_title.font.size = Pt(11)
        r_title.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
        
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
    return p

def add_table_data(doc, headers, rows_data, col_widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    
    # Header Row
    hdr_cells = table.rows[0].cells
    cant_split_row(table.rows[0])
    repeat_header_row(table.rows[0])
    
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "1E293B") # Dark Slate Header
        set_cell_margins(hdr_cells[i], top=140, bottom=140, left=150, right=150)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        for r in p.runs:
            r.bold = True
            r.font.name = 'Arial'
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            
    # Data Rows
    for r_idx, row_values in enumerate(rows_data):
        row = table.add_row()
        cant_split_row(row)
        bg_color = "F8FAFC" if (r_idx % 2 == 1) else "FFFFFF"
        for c_idx, val in enumerate(row_values):
            cell = row.cells[c_idx]
            cell.text = str(val)
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            
            # Align center for short IDs/codes
            if c_idx == 0 or len(str(val)) <= 6 or str(val) in ['Alta', 'Média', 'Baixa', 'Sim', 'Não', 'Aprovado']:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                
            for r in p.runs:
                r.font.name = 'Arial'
                r.font.size = Pt(9)
                r.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
                
    # Widths
    if col_widths and len(col_widths) == len(headers):
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Cm(w)
                
    # Retorna a tabela sem adicionar parágrafos em branco que geram páginas vazias
    return table
