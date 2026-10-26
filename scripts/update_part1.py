import os
import sys
import docx
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx_helper import (
    get_base_doc_path, get_target_doc_path, get_workspace_target_doc_path,
    add_heading_1, add_heading_2, add_heading_3, add_body_p, add_bullet_p
)

def run_part1():
    target_path = get_target_doc_path()
    workspace_target = get_workspace_target_doc_path()
    
    print(f"[PARTE 1] Carregando documento pré-textual limpo: {target_path}")
    doc = docx.Document(target_path)

    # Início do Conteúdo Textual Estruturado ABNT
    print("[PARTE 1] Injetando Capítulo 1: Introdução, Contextualização e Justificativa...")
    
    # Quebra de página antes do Capítulo 1
    doc.add_page_break()
    
    add_heading_1(doc, "1. INTRODUÇÃO")
    
    # Texto oficial base obrigatório expandido
    add_body_p(
        doc,
        "Na atual era da informação, as boas práticas de segurança tornaram-se indispensáveis, "
        "uma vez que a tecnologia passou a integrar de forma definitiva praticamente todas as atividades "
        "profissionais. A crescente digitalização dos processos, o uso intensivo de sistemas informatizados, "
        "a adoção consolidada do trabalho remoto e híbrido e a troca contínua de dados e informações corporativas "
        "ampliaram significativamente a superfície de ataque e a exposição das organizações a diversos riscos cibernéticos, "
        "como vazamento de dados confidenciais, fraudes eletrônicas, ataques de ransomware, tentativas sofisticadas de "
        "engenharia social e o uso inadequado de recursos tecnológicos corporativos."
    )
    
    add_body_p(
        doc,
        "Observa-se que grande parte desses incidentes não decorre exclusivamente de falhas técnicas nos sistemas, "
        "mas, sobretudo, de comportamentos humanos inadequados, falta de atenção, desconhecimento ou ausência de uma cultura "
        "voltada à segurança da informação, o que evidencia a relevância crucial da conscientização e da educação digital no "
        "ambiente organizacional. Segundo relatórios globais de inteligência em ameaças cibernéticas (como o Verizon Data Breach "
        "Investigations Report), mais de 74% das violações de segurança corporativa envolvem diretamente o elemento humano, "
        "seja por meio de ataques de phishing, uso de credenciais fracas ou reutilizadas, erros operacionais de configuração "
        "ou interação negligente com arquivos maliciosos recebidos por canais de comunicação."
    )
    
    add_body_p(
        doc,
        "A adoção de boas práticas de segurança da informação pressupõe o uso responsável e ético da tecnologia, a proteção "
        "de dados sensíveis e estratégicos, o cumprimento rigoroso das políticas internas e das normas institucionais, bem "
        "como o desenvolvimento de um senso crítico diante de e-mails, links, aplicativos e solicitações que possam representar "
        "potenciais ameaças. Nesse contexto, a segurança da informação não deve ser compreendida apenas como um conjunto estático "
        "de medidas técnicas de TI, mas como um processo contínuo de orientação, prevenção e responsabilidade compartilhada "
        "entre todos os colaboradores da empresa."
    )
    
    add_heading_2(doc, "1.1 A Intersecção entre Cibersegurança e Bem-Estar Digital")
    
    add_body_p(
        doc,
        "Além dos aspectos estritamente técnicos, a segurança da informação está diretamente relacionada ao conceito de "
        "bem-estar digital no ambiente de trabalho. O uso excessivo de dispositivos tecnológicos, a conectividade permanente "
        "(cultura do 'always-on') e a pressão contínua por respostas imediatas geram sobrecarga cognitiva, estresse ocupacional "
        "e dispersão da atenção. Esse estado de fadiga tecnológica reduz a capacidade de discernimento dos profissionais, "
        "tornando-os significativamente mais vulneráveis a táticas de engenharia social baseadas em senso de urgência e medo."
    )
    
    add_body_p(
        doc,
        "A ausência de limites no uso da tecnologia e a proliferação desordenada de alertas em ferramentas corporativas criam "
        "a chamada 'fadiga de alertas' (alert fatigue), na qual avisos legítimos de risco passam a ser ignorados ou descartados "
        "automaticamente pelos usuários. A exaustão digital compromete tanto a saúde física e mental dos profissionais quanto a "
        "postura defensiva da organização, reforçando a necessidade imperativa de práticas e ferramentas que promovam um "
        "equilíbrio saudável entre produtividade, segurança cibernética e qualidade de vida no ambiente corporativo."
    )
    
    add_heading_2(doc, "1.2 Justificativa e Impactos Organizacionais")
    
    add_body_p(
        doc,
        "Quando a organização investe de forma consistente na construção de uma cultura de segurança da informação aliada "
        "ao uso consciente e equilibrado da tecnologia, os benefícios extrapolam amplamente a proteção pura de dados e sistemas. "
        "Há uma redução substancial de incidentes de segurança e de custos operacionais com recuperação de desastres e remediação "
        "de violações, maior eficiência e confiabilidade nos processos digitais, aumento sustentável da produtividade, "
        "fortalecimento do clima organizacional e maior engajamento dos colaboradores."
    )
    
    add_body_p(
        doc,
        "Adicionalmente, esse compromisso contribui de forma decisiva para o fortalecimento da imagem institucional e para "
        "o aumento da confiança de clientes, fornecedores, parceiros comerciais e órgãos reguladores, atendendo plenamente "
        "às exigências da Lei Geral de Proteção de Dados Pessoais (LGPD - Lei nº 13.709/2018) e às normas internacionais da "
        "família ISO/IEC 27000. Sob essa ótica, a segurança da informação deixa de ser um centro de custos para atuar como "
        "um vetor estratégico de sustentabilidade e competitividade no mercado contemporâneo."
    )
    
    add_heading_2(doc, "1.3 Problema de Pesquisa")
    
    add_body_p(
        doc,
        "Diante do cenário exposto, formula-se o problema central que norteia esta investigação científica e o desenvolvimento "
        "tecnológico aplicado:"
    )
    
    add_body_p(
        doc,
        "Como a concepção e a implementação de uma plataforma tecnológica integrada — que alie gerenciamento criptográfico de "
        "credenciais sob arquitetura Zero-Knowledge, túnel VPN auditável de alto desempenho, detecção heurística de phishing e "
        "módulos de microaprendizagem gamificados — pode mitigar o risco decorrente de falhas humanas e promover simultaneamente "
        "o bem-estar digital e a conformidade regulatória nas organizações?",
        bold_prefix="Questão Norteadora: "
    )
    
    add_heading_2(doc, "1.4 Objetivos da Pesquisa")
    
    add_body_p(doc, "Para responder à problemática proposta, estabelecem-se os seguintes objetivos:", bold_prefix="Objetivo Geral: ")
    add_body_p(
        doc,
        "Desenvolver, validar e documentar a solução de software 'SeguranApp' (RiskZero), composta por aplicativo híbrido "
        "multiplataforma (Android/Web), extensão de navegador e infraestrutura de rede segura WireGuard, destinada a fortalecer "
        "a governança de dados, proteger os colaboradores contra ameaças cibernéticas cotidianas e fomentar hábitos saudáveis "
        "de uso da tecnologia no contexto corporativo."
    )
    
    add_body_p(doc, "Constituem objetivos específicos deste projeto de conclusão de curso:", bold_prefix="Objetivos Específicos:\n")
    add_bullet_p(doc, " Mapear as principais vulnerabilidades associadas ao comportamento humano no ambiente profissional moderno, incluindo phishing, reaproveitamento de senhas, conexões em redes desprotegidas e fadiga informacional;", bold_title="a)")
    add_bullet_p(doc, " Projetar e implementar um cofre de credenciais digital baseado em criptografia simétrica AES-GCM de 256 bits e derivação de chaves por PBKDF2/Argon2id sob a premissa de Conhecimento Zero (Zero-Knowledge Architecture);", bold_title="b)")
    add_bullet_p(doc, " Integrar um cliente e servidor de Virtual Private Network (VPN) utilizando o protocolo WireGuard de última geração em curva elíptica Curve25519, garantindo sigilo em conexões remotas e Wi-Fi público;", bold_title="c)")
    add_bullet_p(doc, " Desenvolver um motor heurístico em tempo real capaz de detectar links fraudulentos, domínios de phishing e artefatos de engenharia social, atribuindo pontuações de risco transparentes;", bold_title="d)")
    add_bullet_p(doc, " Criar módulos interativos de educação digital e quiz gamificado para capacitação continuada dos colaboradores sem gerar sobrecarga cognitiva;", bold_title="e)")
    add_bullet_p(doc, " Estabelecer mecanismos de hardening nativo no dispositivo móvel, tais como bloqueio de capturas de tela (FLAG_SECURE), limpeza programada da área de transferência e geração de logs para conformidade com a LGPD;", bold_title="f)")
    add_bullet_p(doc, " Validar a eficácia e a usabilidade da solução por meio de baterias de testes automatizados unitários, testes de ponta a ponta (E2E) e estudos de viabilidade mercadológica B2B/B2C.", bold_title="g)")
    
    add_heading_2(doc, "1.5 Estrutura da Monografia")
    
    add_body_p(
        doc,
        "Este Trabalho de Conclusão de Curso encontra-se estruturado em capítulos lógicos e interdependentes: "
        "O Capítulo 1 apresentou a introdução, o contexto da era digital, o fator humano, a justificação e os objetivos. "
        "O Capítulo 2 consolida a fundamentação teórica sobre segurança da informação, normas ISO, LGPD, engenharia social, "
        "bem-estar digital e criptografia. "
        "O Capítulo 3 descreve a metodologia científica e a gestão ágil da equipe (Scrum/Kanban) com o Backlog e Cronograma. "
        "O Capítulo 4 formaliza a especificação exaustiva de Requisitos Funcionais e Não Funcionais. "
        "O Capítulo 5 detalha a arquitetura do sistema e a stack tecnológica escolhida. "
        "O Capítulo 6 expõe a modelagem do banco de dados relacional e NoSQL, incluindo o Dicionário de Dados completo. "
        "O Capítulo 7 apresenta a prototipação UI/UX, o protótipo navegável Figma e o fluxo das telas. "
        "O Capítulo 8 consolida as rotinas de testes de software, a matriz de casos de teste e o relatório de incidentes. "
        "O Capítulo 9 formaliza o plano de negócios e a viabilidade mercadológica. "
        "Por fim, o Capítulo 10 tece as considerações finais, seguido pelas referências bibliográficas normatizadas."
    )
    
    # Capítulo 2: Fundamentação Teórica
    print("[PARTE 1] Injetando Capítulo 2: Fundamentação Teórica Aprofundada...")
    doc.add_page_break()
    
    add_heading_1(doc, "2. FUNDAMENTAÇÃO TEÓRICA")
    
    add_body_p(
        doc,
        "A fundamentação teórica deste trabalho apoia-se em literatura acadêmica consagrada, normas internacionais "
        "de padronização e legislação vigente nas áreas de segurança da informação, computação aplicada, psicologia "
        "comportamental no trabalho e ergonomia de sistemas. Como sustentam Whitman e Mattord (2019) e Stallings (2015), "
        "a segurança computacional eficaz transcende a implementação de firewalls e antivírus, exigindo a harmonia perfeita "
        "entre três pilares indissociáveis: pessoas, processos e tecnologia."
    )
    
    add_heading_2(doc, "2.1 Os Pilares da Segurança da Informação (Tripé CID e Extensões)")
    
    add_body_p(
        doc,
        "Classicamente, a segurança da informação estrutura-se sobre o tripé denominado CID (Confidencialidade, Integridade "
        "e Disponibilidade), expandido pela moderna engenharia de software para incorporar a Autenticidade e a Irretratabilidade "
        "(Não-Repúdio), conforme preconizado pela norma ABNT NBR ISO/IEC 27001:2022:"
    )
    
    add_bullet_p(doc, " Propriedade que garante que a informação não seja disponibilizada ou divulgada a indivíduos, entidades ou processos não autorizados. No SeguranApp, a confidencialidade é assegurada mediante o uso de criptografia simétrica de ponta a ponta (AES-GCM de 256 bits), onde apenas o proprietário possui a chave de decifração em sua memória volátil.", bold_title="Confidencialidade (Confidentiality):")
    add_bullet_p(doc, " Salvaguarda da exatidão e da completude da informação e dos métodos de processamento. Protege os dados contra alterações indevidas, exclusões acidentais ou adulterações maliciosas, validada criptograficamente através de assinaturas e autenticação de mensagens (Authentication Tags no AES-GCM).", bold_title="Integridade (Integrity):")
    add_bullet_p(doc, " Garantia de que os usuários autorizados tenham acesso imediato e confiável à informação e aos ativos correspondentes sempre que demandado. No projeto, a alta disponibilidade é viabilizada por arquitetura em nuvem distribuída e suporte a operação offline com sincronização assíncrona.", bold_title="Disponibilidade (Availability):")
    add_bullet_p(doc, " Garantia de que a identidade alegada por um usuário, dispositivo ou serviço corresponda efetivamente à sua identidade real, materializada no sistema por autenticação biométrica (WebAuthn) e senhas mestras salteadas.", bold_title="Autenticidade (Authenticity):")
    add_bullet_p(doc, " Capacidade de provar a autoria e a ocorrência de determinada ação ou transação digital, impedindo que o emissor negue sua autoria, suportada por registros imutáveis de trilhas de auditoria (Audit Logs).", bold_title="Não-Repúdio (Non-repudiation):")
    
    add_heading_2(doc, "2.2 Engenharia Social, Vetores de Ataque e a Falha Humana")
    
    add_body_p(
        doc,
        "Conforme conceitua Mitnick e Simon (2003), a engenharia social consiste na arte de manipular indivíduos "
        "psicologicamente para que executem ações precipitadas ou divulguem informações sigilosas. Ao contrário dos ataques "
        "voltados à exploração de falhas de código (*software exploits*), a engenharia social explora vulnerabilidades inerentes "
        "à cognição humana, tais como a empatia, a obediência à autoridade, o medo de punição, a curiosidade e o senso de urgência."
    )
    
    add_body_p(
        doc,
        "Entre as modalidades mais frequentes e danosas no contexto corporativo, destacam-se:",
        bold_prefix="Principais Vetores de Ataque Comportamental:\n"
    )
    
    add_bullet_p(doc, " Disseminação massiva de e-mails, mensagens instantâneas e páginas falsas que imitam instituições bancárias, portais de recursos humanos ou provedores de nuvem para captura de credenciais.", bold_title="Phishing Tradicional e Spear Phishing:")
    add_bullet_p(doc, " Ataques de phishing altamente customizados direcionados a colaboradores específicos, utilizando dados coletados em redes sociais profissionais (como LinkedIn) para conferir verossimilhança à solicitação.", bold_title="Engenharia Social Dirigida:")
    add_bullet_p(doc, " Fraudes perpetradas por chamadas de voz telefônicas (Vishing) e mensagens SMS ou WhatsApp (Smishing), simulando suporte técnico emergencial ou supostos bloqueios de contas bancárias e PIX.", bold_title="Vishing e Smishing:")
    add_bullet_p(doc, " Testes automatizados executados por botnets que utilizam listas massivas de e-mails e senhas vazadas em outros serviços na internet para tentar invadir contas corporativas, explorando o hábito perigoso de reaproveitamento de senhas.", bold_title="Credential Stuffing:")
    add_bullet_p(doc, " Interceptação clandestina do tráfego de rede em roteadores vulneráveis ou pontos de acesso Wi-Fi públicos desprotegidos, permitindo a captura de tokens de sessão e credenciais em texto claro.", bold_title="Man-in-the-Middle (MitM):")
    add_bullet_p(doc, " Registro de domínios visualmente idênticos a marcas renomadas substituindo caracteres do alfabeto latino por glifos cirílicos ou utilizando técnicas de Punycode para induzir o usuário ao erro óptico.", bold_title="Ataques Homógrafos e Spoofing de URLs:")
    
    add_heading_2(doc, "2.3 Marco Regulatório: Lei Geral de Proteção de Dados (LGPD - Lei nº 13.709/2018)")
    
    add_body_p(
        doc,
        "Promulgada em agosto de 2018 e em pleno vigor no Brasil, a Lei nº 13.709/2018 estabelece o regime jurídico para "
        "o tratamento de dados pessoais em meios físicos e digitais por pessoas naturais e jurídicas de direito público ou privado. "
        "Inspirada diretamente no Regulamento Geral sobre a Proteção de Dados da União Europeia (GDPR 2016/679), a LGPD "
        "reconfigurou o tratamento de dados no ambiente corporativo brasileiro."
    )
    
    add_body_p(
        doc,
        "O artigo 6º da legislação estabelece os dez princípios fundamentais que devem reger qualquer ciclo de vida do dado, "
        "destacando-se para o escopo desta pesquisa: o Princípio da Segurança (adoção de medidas técnicas e administrativas aptas "
        "a proteger os dados de acessos não autorizados e de situações acidentais ou ilícitas de destruição, perda, alteração ou difusão), "
        "o Princípio da Prevenção (adoção de medidas prévias para evitar danos) e o Princípio da Responsabilização e Prestação de Contas "
        "(capacidade do agente de tratamento comprovar a eficácia e conformidade de suas medidas)."
    )
    
    add_body_p(
        doc,
        "O não cumprimento das diretrizes estabelecidas na lei sujeita as organizações infratoras a severas penalidades "
        "administrativas aplicadas pela Autoridade Nacional de Proteção de Dados (ANPD), que variam desde advertências públicas "
        "até multas de até 2% do faturamento da pessoa jurídica (limitadas a 50 milhões de reais por infração), além do bloqueio "
        "e eliminação dos dados tratados de forma irregular, sem prejuízo de reparações civis por danos morais individuais e coletivos. "
        "Nesse sentido, ferramentas que implementam cifragem local e auditoria sem custódia de dados (Zero-Knowledge) atuam como "
        "salvaguardas diretas de conformidade legal."
    )
    
    add_heading_2(doc, "2.4 Normas Internacionais: Família ABNT NBR ISO/IEC 27000")
    
    add_body_p(
        doc,
        "A série de normas ISO/IEC 27000 representa o padrão ouro global para concepção, implementação, operação, monitoramento "
        "e melhoria contínua de um Sistema de Gestão de Segurança da Informação (SGSI):"
    )
    
    add_bullet_p(doc, " Especifica os requisitos formais para estabelecer, implementar, manter e aprimorar continuamente um SGSI, fundamentado na gestão de riscos cibernéticos e no comprometimento da alta direção.", bold_title="ABNT NBR ISO/IEC 27001:2022:")
    add_bullet_p(doc, " Funciona como guia prático de referência, fornecendo um catálogo de 93 controles de segurança agrupados em 4 temas centrais: Controles Organizacionais (37 controles), Controles de Pessoas (8 controles, enfatizando conscientização, triagem e responsabilidades remotas), Controles Físicos (14 controles) e Controles Tecnológicos (34 controles, abrangendo criptografia, autenticação segura, mascaramento de dados e proteção contra malware).", bold_title="ABNT NBR ISO/IEC 27002:2022:")
    
    add_heading_2(doc, "2.5 Bem-Estar Digital, Ergonomia Cognitiva e Sobrecarga Tecnológica")
    
    add_body_p(
        doc,
        "O conceito de Bem-Estar Digital (Digital Wellbeing) refere-se ao estado de harmonia e equilíbrio psicológico "
        "entre o indivíduo e as tecnologias digitais com as quais interage profissionalmente. Segundo a Teoria da Carga "
        "Cognitiva formulada por Sweller (1988), a memória de trabalho humana possui capacidade estritamente limitada para "
        "processar novas informações simultaneamente."
    )
    
    add_body_p(
        doc,
        "Quando interfaces corporativas são mal projetadas, gerando sobrecarga de notificações, múltiplos alertas ambíguos "
        "e excesso de etapas burocráticas para execução de procedimentos rotineiros, ocorre a sobrecarga cognitiva extrínseca "
        "e o estresse tecnológico (*technostress*). Conforme demonstram pesquisas em fatores humanos aplicados à segurança, "
        "o estresse mental crônico e a fragmentação da atenção reduzem a atividade do córtex pré-frontal, induzindo o "
        "colaborador a operar no chamado 'Modo Piloto Automático'."
    )
    
    add_body_p(
        doc,
        "Nesse estado de fadiga, o colaborador clica precipitadamente em anexos fraudulentos, desativa mecanismos de defesa "
        "para poupar tempo ou anota senhas confidenciais em papéis sobre a mesa. Portanto, a simplificação da interface, o "
        "uso de paletas visuais escuras e limpas (*Dark Mode*) que reduzem a fadiga ocular, o preenchimento automático "
        "sem fricção e o fornecimento de feedback educativo claro e construtivo (ao invés de mensagens de erro punitivas) "
        "constituem requisitos funcionais indispensáveis para uma segurança digital sustentável."
    )
    
    add_heading_2(doc, "2.6 Arquitetura Criptográfica Contemporânea e Redes Seguras")
    
    add_body_p(
        doc,
        "Para transformar os preceitos teóricos em soluções concretas e auditáveis, este trabalho fundamenta-se nas mais "
        "modernas primitivas e protocolos de segurança digital:",
        bold_prefix="Fundamentos Tecnológicos Adotados:\n"
    )
    
    add_bullet_p(doc, " Cifra simétrica de bloco autenticado operando em modo Galois/Counter. Fornece simultaneamente confidencialidade e integridade garantida por tag de autenticação GMAC (Galois Message Authentication Code), imune a ataques de texto cifrado escolhido (IND-CCA2) e altamente veloz por dispensar computações adicionais de HMAC.", bold_title="Advanced Encryption Standard em Modo GCM (AES-GCM 256):")
    add_bullet_p(doc, " Funções criptográficas de estiramento de chave (Key Derivation Functions). O PBKDF2 aplica 600.000 iterações de HMAC-SHA256 com salteamento aleatório exclusivo de 16 bytes para frustrar ataques de dicionário e tabelas rainbow, enquanto o algoritmo Argon2id (vencedor do Password Hashing Competition) garante resistência máxima contra ataques paralelos acelerados por GPUs e ASICs.", bold_title="Derivação de Chaves (PBKDF2 e Argon2id):")
    add_bullet_p(doc, " Mecanismo no qual a senha mestra do usuário não cifra diretamente os milhares de itens do cofre. Em vez disso, a chave derivada atua como uma Chave de Cifragem de Chave (KEK), que cifra uma única Chave de Cifragem de Dados (DEK) de 256 bits gerada aleatoriamente. Isso viabiliza a alteração da senha mestra instantaneamente com a re-cifragem de apenas uma chave, sem tocar na base inteira de credenciais.", bold_title="Cifragem por Envelope (Envelope Encryption):")
    add_bullet_p(doc, " Algoritmo padronizado pela IETF para autenticação multifator (MFA) baseada em tempo. Calcula códigos numéricos de seis dígitos derivados de um segredo compartilhado em Base32 e do carimbo de tempo UNIX atual em janelas estritas de 30 segundos, eliminando a dependência de SMS vulneráveis a SIM Swap.", bold_title="Time-based One-Time Password (TOTP - RFC 6238):")
    add_bullet_p(doc, " Protocolo moderno de VPN que opera diretamente no espaço de kernel com base de código enxuta (~4.000 linhas de código C contra mais de 100.000 do OpenVPN). Utiliza criptografia assimétrica de curva elíptica Curve25519 para troca de chaves Diffie-Hellman, ChaCha20 para cifragem simétrica, Poly1305 para autenticação e BLAKE2s para hashing, proporcionando latências insignificantes e consumo reduzido de bateria em dispositivos móveis.", bold_title="Protocolo WireGuard (VPN de Última Geração):")

    # Salvar em TCC_Documentacao_Completa.docx
    doc.save(target_path)
    # Também salvar na raiz do workspace
    try:
        doc.save(workspace_target)
    except Exception as e:
        print(f"Aviso ao salvar na raiz do workspace: {e}")
        
    print(f"[PARTE 1 CONCLUÍDA COM SUCESSO!]")
    print(f"Arquivo gerado: {target_path}")
    print(f"Total de parágrafos no documento após Parte 1: {len(doc.paragraphs)}")

if __name__ == "__main__":
    run_part1()
