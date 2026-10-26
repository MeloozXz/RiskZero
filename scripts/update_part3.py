import os
import sys
import docx
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx_helper import (
    get_target_doc_path, get_workspace_target_doc_path,
    add_heading_1, add_heading_2, add_heading_3, add_body_p, add_bullet_p, add_table_data
)

def generate_mermaid_file():
    mermaid_content = """# Diagramas do Projeto SeguranApp (RiskZero) em Mermaid.js

Este documento reúne a especificação visual e lógica dos diagramas do Trabalho de Conclusão de Curso em sintaxe padrão Mermaid.js. Podem ser visualizados no GitHub, VS Code ou renderizados em PNG/SVG através do [Mermaid Live Editor](https://mermaid.live).

---

## 1. Diagrama de Arquitetura do Sistema e Fluxo de Dados

```mermaid
graph TD
    subgraph "Camada de Apresentação (Cliente)"
        A1["📱 App Mobile Android<br/>(Capacitor 6 + Ionic React)"]
        A2["💻 Web Application<br/>(React 19 + Tailwind v4)"]
        A3["🧩 Extensão de Navegador<br/>(Chrome Manifest V3)"]
    end

    subgraph "Camada de Segurança Local (Client-Side)"
        B1["🔐 Web Crypto API<br/>(AES-GCM 256 / PBKDF2)"]
        B2["🛡️ Android Keystore & Biometrics<br/>(WebAuthn / BiometricPrompt)"]
        B3["⏱️ RFC 6238 TOTP Engine<br/>(Gerador 2FA Local)"]
        B4["🔍 Motor Heurístico Anti-Phishing<br/>(Threat Score Calculator)"]
        B5["🔒 Privacy Screen & Clipboard<br/>(FLAG_SECURE / Auto-Clear 30s)"]
    end

    subgraph "Camada de Rede Segura e Túneis"
        C1["🌐 VPS Linux Ubuntu 22.04 LTS<br/>(WireGuard Server / UDP 51820)"]
        C2["☁️ Cloudflare WARP Edge Network<br/>(Anycast WireGuard Route)"]
        C3["🛡️ Resolvedores DNS Criptografados<br/>(1.1.1.1 / Quad9 Anti-Leak)"]
    end

    subgraph "Serviços Externos e Nuvem Zero-Knowledge"
        D1["☁️ Firebase Cloud Firestore<br/>(Armazenamento Cifrado por UID)"]
        D2["🔑 Firebase Authentication<br/>(Gestão de Sessão do Usuário)"]
        D3["🔎 HaveIBeenPwned API<br/>(k-Anonymity SHA-1 Prefix Lookup)"]
    end

    %% Conexões Cliente -> Segurança Local
    A1 --> B1
    A1 --> B2
    A1 --> B3
    A1 --> B4
    A1 --> B5
    A2 --> B1
    A2 --> B4
    A3 --> B1
    A3 --> B4

    %% Conexões Segurança Local -> Rede & Nuvem
    B1 -- "Envelopes Cifrados (Ciphertext + IV)" --> D1
    A1 -- "Autenticação de Usuário" --> D2
    A1 -- "Túnel WireGuard Curve25519" --> C1
    A1 -- "Alternativa WARP WireGuard" --> C2
    C1 --> C3
    B4 -- "Consulta k-Anonymity (5 primeiros chars SHA-1)" --> D3
```

---

## 2. Diagrama Entidade-Relacionamento (DER Lógico)

```mermaid
erDiagram
    USUARIOS_PERFIL ||--|| COFRE_ENVELOPE : "possui"
    USUARIOS_PERFIL ||--o{ LOGS_AUDITORIA : "gera"
    USUARIOS_PERFIL ||--o{ SESSOES_ATIVAS : "mantém"
    USUARIOS_PERFIL ||--o| ASSINATURAS_PLANOS : "assina"
    COFRE_ENVELOPE ||--o{ COFRE_ITENS : "contém"
    COFRE_ITENS ||--o{ HISTORICO_SENHAS : "registra"
    USUARIOS_PERFIL ||--o{ SERVIDORES_VPN : "conecta"

    USUARIOS_PERFIL {
        string id PK "Identificador único (UID Firebase)"
        string username "E-mail ou nome de login"
        string salt "Salt aleatório CSPRNG de 16 bytes"
        string kdfAlgorithm "Algoritmo KDF (pbkdf2 ou argon2id)"
        string passwordHash "Hash SHA-256 (senha mestra + salt)"
        string pinHash "Hash SHA-256 do PIN local de 6 dígitos"
        string encryptedTotpSecret "Segredo TOTP cifrado com chave mestra"
        boolean totpEnabled "Status de ativação do 2FA"
        string createdAt "Carimbo de data/hora de criação"
    }

    COFRE_ENVELOPE {
        string id PK "Identificador do cofre"
        string userId FK "Chave estrangeira para USUARIOS_PERFIL"
        int version "Versão do esquema do cofre"
        string encryptedDekCiphertext "Chave DEK (256 bits) cifrada pela KEK"
        string encryptedDekIv "Vetor de inicialização do envelope DEK"
        string updatedAt "Data da última sincronização"
    }

    COFRE_ITENS {
        string id PK "Identificador único do item de credencial"
        string envelopeId FK "Chave estrangeira para COFRE_ENVELOPE"
        string title "Título ou nome do serviço cadastrado"
        string username "Nome de usuário ou e-mail de acesso"
        string encryptedPasswordCiphertext "Senha cifrada com DEK (AES-GCM 256)"
        string encryptedPasswordIv "Vetor de inicialização do item"
        string url "Endereço web ou host do serviço"
        string category "LOGIN, CARD, NOTE, IDENTITY, OTHER"
        string encryptedNotesCiphertext "Notas adicionais cifradas"
        string encryptedNotesIv "Vetor de inicialização das notas"
        string encryptedTotpSecret "Segredo TOTP individual cifrado"
        string updatedAt "Carimbo de data da última modificação"
    }

    HISTORICO_SENHAS {
        string id PK "Identificador da entrada de histórico"
        string itemId FK "Chave estrangeira para COFRE_ITENS"
        string username "Usuário na época da alteração"
        string encryptedPasswordCiphertext "Senha anterior cifrada"
        string encryptedPasswordIv "Vetor de inicialização da senha histórica"
        string updatedAt "Data da substituição da senha"
    }

    LOGS_AUDITORIA {
        string id PK "Identificador do registro de auditoria"
        string userId FK "Chave estrangeira para USUARIOS_PERFIL"
        string timestamp "Carimbo ISO 8601 do evento"
        string action "LOGIN_SUCCESS, VPN_CONNECTED, AUTO_LOCK..."
        string details "Descrição textual do evento de segurança"
    }

    SERVIDORES_VPN {
        string id PK "Identificador do nó de rede WireGuard"
        string name "Nome descritivo (ex: VPS Brasil 01)"
        string country "País do servidor"
        string city "Cidade de localização"
        string flag "Emoji de bandeira do país"
        string endpoint "Endereço IP público ou hostname do servidor"
        int port "Porta UDP de escuta (51820)"
        string serverPublicKey "Chave pública Curve25519 (Base64)"
        string clientAddress "IP do cliente na sub-rede VPN (10.14.0.2/32)"
        string status "AVAILABLE, UNAVAILABLE, MAINTENANCE"
        boolean isCustom "Indica servidor particular cadastrado"
        int ping "Latência média em milissegundos"
    }

    ASSINATURAS_PLANOS {
        string id PK "Identificador do contrato de assinatura"
        string userId FK "Chave estrangeira para USUARIOS_PERFIL"
        string tier "Nível do plano: FREE, PRO, ULTRA"
        string billingCycle "MONTHLY ou ANNUAL"
        string status "ACTIVE, EXPIRED, TRIAL"
        string startedAt "Data de início da vigência"
        string expiresAt "Data de expiração ou null (vitalício)"
        boolean autoRenew "Renovação automática ativada"
        string paymentMethod "PIX ou CREDIT_CARD"
    }

    SESSOES_ATIVAS {
        string id PK "Identificador da sessão ativa"
        string userId FK "Chave estrangeira para USUARIOS_PERFIL"
        string deviceName "Nome do smartphone ou navegador"
        string lastActive "Carimbo de última interação"
        string ipPlaceholder "Indicação de IP protegido ou VPN"
    }
```

---

## 3. Modelo Conceitual (MER Conceitual)

```mermaid
classDiagram
    class USUARIO {
        +ID_Usuario
        +Nome_Acesso
        +Credenciais_Mestras
        +Fator_Biometrico
        +Configuracao_2FA
    }

    class COFRE_SEGURANCA {
        +ID_Cofre
        +Chave_DEK_Cifrada
        +Versao_Estrutural
    }

    class CREDENCIAL {
        +ID_Item
        +Servico
        +Login
        +Senha_Cifrada
        +URL_Destino
        +Categoria
        +Notas_Cifradas
    }

    class HISTORICO_SENHA {
        +ID_Historico
        +Senha_Antiga_Cifrada
        +Data_Alteracao
    }

    class TRILHA_AUDITORIA {
        +ID_Log
        +Data_Hora
        +Tipo_Acao
        +Detalhes_Evento
    }

    class SERVIDOR_VPN {
        +ID_Servidor
        +Host_IP
        +Porta_UDP
        +Chave_Publica_WireGuard
        +Localidade
    }

    class PLANO_ASSINATURA {
        +ID_Plano
        +Nivel_Tier
        +Status_Vigencia
        +Ciclo_Cobranca
    }

    USUARIO "1" -- "1" COFRE_SEGURANCA : gerencia
    COFRE_SEGURANCA "1" -- "*" CREDENCIAL : encapsula
    CREDENCIAL "1" -- "*" HISTORICO_SENHA : rastreia
    USUARIO "1" -- "*" TRILHA_AUDITORIA : audita
    USUARIO "1" -- "*" SERVIDOR_VPN : conecta
    USUARIO "1" -- "1" PLANO_ASSINATURA : contrata
```

---

## 4. Gráfico de Gantt do Cronograma de Sprints

```mermaid
gantt
    title Cronograma de Desenvolvimento Ágil — SeguranApp (RiskZero)
    dateFormat  YYYY-MM-DD
    section Sprint 1: Fundação & UI/UX
    Levantamento de Requisitos e Arquitetura Zero-Knowledge :s1_1, 2026-02-02, 7d
    Prototipação no Figma e Configuração do Repositório :s1_2, 2026-02-09, 7d
    section Sprint 2: Criptografia & Cofre
    Implementação Web Crypto (AES-GCM 256 / PBKDF2) :s2_1, 2026-02-16, 7d
    Desenvolvimento do Cofre e Biometria / PIN :s2_2, 2026-02-23, 7d
    section Sprint 3: Rede & WireGuard VPN
    Provisionamento VPS Linux e Configuração WireGuard :s3_1, 2026-03-02, 7d
    Cliente VPN Android, Chaves Curve25519 e Cloudflare WARP :s3_2, 2026-03-09, 7d
    section Sprint 4: Inteligência & Hardening
    Motor Heurístico Anti-Phishing e Threat Score :s4_1, 2026-03-16, 7d
    Auditoria k-Anonymity HIBP e Privacy Screen FLAG_SECURE :s4_2, 2026-03-23, 7d
    section Sprint 5: Educação & Extensão
    Central de Educação Digital (13 tópicos) e Quiz :s5_1, 2026-03-30, 7d
    Extensão Chrome MV3 e Sistema de Assinaturas SaaS :s5_2, 2026-04-06, 7d
    section Sprint 6: Testes & Monografia
    Testes Automatizados Vitest & Cypress E2E :s6_1, 2026-04-13, 7d
    Auditoria Final e Consolidação Monográfica ABNT :s6_2, 2026-04-20, 7d
```
"""
    docs_dir = os.path.join(os.path.dirname(__file__), "..", "docs")
    os.makedirs(docs_dir, exist_ok=True)
    out_file = os.path.join(docs_dir, "DIAGRAMAS_MERMAID.md")
    with open(out_file, "w", encoding="utf-8") as f:
        f.write(mermaid_content)
    print(f"[PARTE 3] Arquivo Mermaid gerado com sucesso em: {out_file}")

    # Também salvar na raiz workspace docs se diferente
    workspace_docs = os.path.join("C:\\Users\\Aluno\\Downloads\\seguran-aapp-main (2)", "docs")
    os.makedirs(workspace_docs, exist_ok=True)
    with open(os.path.join(workspace_docs, "DIAGRAMAS_MERMAID.md"), "w", encoding="utf-8") as f:
        f.write(mermaid_content)

def run_part3():
    target_path = get_target_doc_path()
    workspace_target = get_workspace_target_doc_path()
    
    # 1. Gerar arquivo docs/DIAGRAMAS_MERMAID.md
    generate_mermaid_file()
    
    print(f"[PARTE 3] Carregando documento gerado pela Parte 2: {target_path}")
    doc = docx.Document(target_path)
    
    # Quebra de página antes do Capítulo 5
    doc.add_page_break()
    
    # =========================================================================
    # CAPÍTULO 5: ARQUITETURA DO SISTEMA E STACK TECNOLÓGICA
    # =========================================================================
    print("[PARTE 3] Injetando Capítulo 5: Arquitetura do Sistema e Stack Tecnológica...")
    add_heading_1(doc, "5. ARQUITETURA DO SISTEMA E STACK TECNOLÓGICA")
    
    add_heading_2(doc, "5.1 Padrão Arquitetural Adotado: Cliente-Servidor Descentralizado e Zero-Knowledge")
    add_body_p(
        doc,
        "A arquitetura do ecossistema SeguranApp (RiskZero) foi concebida sob o modelo Cliente-Servidor Híbrido Descentralizado, "
        "alicerçado na rigorosa premissa de Conhecimento Zero (Zero-Knowledge Architecture). Em uma arquitetura convencional em nuvem, "
        "o servidor atua como repositório fiduciário dos dados, decifrando ou processando informações em texto claro. "
        "No SeguranApp, a fronteira de confiança (Trust Boundary) é estritamente restrita ao ambiente de execução local do "
        "dispositivo do usuário (smartphone ou navegador)."
    )
    add_body_p(
        doc,
        "O sistema organiza-se em quatro camadas interdependentes e desacopladas:",
        bold_prefix="Camadas da Arquitetura SeguranApp:\n"
    )
    add_bullet_p(doc, " Composta pelo aplicativo móvel híbrido empacotado para Android (via Capacitor 6), pela aplicação web responsiva (Vite + React 19) e pela extensão de navegador Google Chrome (Manifest V3). Responsável pela renderização de alta fidelidade visual, ergonomia cognitiva e captura das ações do usuário.", bold_title="1. Camada de Apresentação e Interação:")
    add_bullet_p(doc, " Opera isoladamente na memória do dispositivo. Executa as operações criptográficas pesadas (AES-GCM de 256 bits, derivação PBKDF2/Argon2id, cálculo de tokens TOTP e rotação de DEK), gerencia o bloqueio de tela (FLAG_SECURE) e sanitiza a área de transferência.", bold_title="2. Camada de Segurança Criptográfica Local:")
    add_bullet_p(doc, " Constituída por Virtual Private Servers (VPS) Linux dedicadas (Ubuntu 22.04 LTS / 24.04 LTS) operando instâncias de alto desempenho do WireGuard na porta UDP 51820, complementadas por nós de borda da rede Anycast Cloudflare WARP e resolvedores DNS anti-vazamento (1.1.1.1 e Quad9).", bold_title="3. Camada de Rede Segura e Tunelamento WireGuard:")
    add_bullet_p(doc, " Formada pelo Firebase Authentication (gestão de identificadores de acesso e tokens de sessão) e Cloud Firestore (banco NoSQL distribuído). O Firestore atua estritamente como um cofre cego de sincronização: recebe e despacha apenas payloads cifrados ('ciphertext', 'iv', 'salt'), sendo tecnicamente incapaz de decifrar o acervo de credenciais.", bold_title="4. Camada de Sincronização em Nuvem (Cofre Cego):")

    add_heading_2(doc, "5.2 Fluxo de Dados Ponta a Ponta (End-to-End Data Flow)")
    add_body_p(
        doc,
        "O fluxo de dados para cadastro, armazenamento e sincronização de uma nova credencial segue um protocolo rigoroso em sete etapas:",
        bold_prefix="Ciclo Operacional de Cifragem e Persistência:\n"
    )
    add_bullet_p(doc, " O usuário insere a Senha Mestra na tela de login. A Web Crypto API executa 600.000 iterações de HMAC-SHA256 combinando a senha ao salt de 16 bytes do usuário, derivando na memória RAM volátil a Chave de Cifragem de Chave (KEK).", bold_title="Etapa 1 (Derivação da Chave na Memória):")
    add_bullet_p(doc, " A KEK decifra o envelope criptográfico que contém a Chave de Cifragem de Dados (DEK de 256 bits). A partir desse momento, a DEK fica disponível temporariamente na sessão local.", bold_title="Etapa 2 (Abertura do Envelope DEK):")
    add_bullet_p(doc, " Ao cadastrar uma credencial (ex: senha corporativa), o gerador criptográfico do sistema (CSPRNG) gera um Vetor de Inicialização (IV) de 96 bits inédito e exclusivo para aquele item.", bold_title="Etapa 3 (Geração de IV Aleatório Exclusivo):")
    add_bullet_p(doc, " O algoritmo AES-GCM processa a senha em texto claro, o IV e a DEK, produzindo o texto cifrado (ciphertext) e a tag de integridade GMAC de 128 bits, serializados em Base64.", bold_title="Etapa 4 (Cifragem Simétrica Autenticada):")
    add_bullet_p(doc, " O item cifrado é gravado imediatamente no armazenamento local persistente seguro do aparelho (Capacitor Preferences), permitindo funcionamento offline irrestrito.", bold_title="Etapa 5 (Persistência Segura Offline):")
    add_bullet_p(doc, " Se a conectividade com a internet estiver ativa e o plano do usuário incluir sincronização em nuvem, o aplicativo envia o payload cifrado via HTTPS/TLS 1.3 para o Firestore sob o caminho 'users/{userId}'. As regras de segurança garantem isolamento multilocatário absoluto.", bold_title="Etapa 6 (Sincronização em Nuvem Cega):")
    add_bullet_p(doc, " A credencial permanece no formulário apenas durante o tempo de edição. Ao sair da tela ou após inatividade do usuário, todas as instâncias em texto claro são destruídas na memória do processo.", bold_title="Etapa 7 (Zeroização de Memória):")

    add_heading_2(doc, "5.3 Justificativa Técnica Minuciosa da Stack Tecnológica")
    add_body_p(
        doc,
        "Cada componente tecnológico da solução foi criteriosamente selecionado com base em benchmarks de mercado, "
        "maturidade de segurança, desempenho e suporte a ecossistemas móveis corporativos modernos:",
        bold_prefix="Matriz de Justificativas Tecnológicas:\n"
    )
    add_bullet_p(doc, " Framework moderno para interfaces híbridas baseado em Web Components padronizados. Proporciona fidelidade visual idêntica à de aplicativos nativos Android e iOS, suporte a temas e transições suaves sem penalizar a memória.", bold_title="Ionic React 8.5 & React 19:")
    add_bullet_p(doc, " A tipagem estática rigorosa do TypeScript eliminou vulnerabilidades comuns de tipagem fraca em tempo de compilação, essenciais para a manipulação de arrays binários (Uint8Array, ArrayBuffer) na criptografia. O Vite substituiu bundlers legados, propiciando build ultrarrápido por Hot Module Replacement (HMR).", bold_title="TypeScript 5.4 & Vite 5:")
    add_bullet_p(doc, " Ponte nativa de última geração que substitui o Apache Cordova com arquitetura moderna orientada a Promises e controle total do projeto Android Studio nativo em Java/Gradle. Viabilizou a integração direta com plugins essenciais de segurança (@capacitor-community/privacy-screen).", bold_title="Capacitor 6 (Android Native Runtime):")
    add_bullet_p(doc, " Motor CSS utilitário que viabilizou o design system ergonômico em Dark Mode, garantindo alto contraste acessível e tempo de carregamento instantâneo por compilação de classes just-in-time.", bold_title="Tailwind CSS v4 & Framer Motion:")
    add_bullet_p(doc, " Renderização de gráficos 3D baseada em WebGL para o escudo de cibersegurança na tela inicial. Transmite feedback visual dinâmico em tempo real sobre o estado de proteção do aparelho, elevando o engajamento do colaborador.", bold_title="Spline 3D (@splinetool/react-spline):")
    add_bullet_p(doc, " Padrão nativo dos navegadores modernos e WebViews (W3C Recommendation). Executa operações criptográficas (AES-GCM, PBKDF2) diretamente em código compilado nativo do sistema operacional (C++ no motor V8/Blink), garantindo desempenho cem vezes superior a bibliotecas JavaScript puras e resistência a ataques de temporização (Side-Channel Timing Attacks).", bold_title="W3C Web Cryptography API:")
    add_bullet_p(doc, " Protocolo moderno de VPN de código aberto operando no espaço de kernel do Linux. Apresenta base de código enxuta (~4.000 linhas), criptografia assimétrica de curva elíptica Curve25519 e cifra ChaCha20-Poly1305, reduzindo a latência em até 40% em relação ao OpenVPN e IPsec.", bold_title="Protocolo WireGuard & VPS Linux:")
    add_bullet_p(doc, " Especificação moderna do Google para extensões de navegador com Service Workers isolados em segundo plano, garantindo que o módulo de preenchimento automático opere sem permissões excessivas e sem degradar o desempenho do navegador corporativo.", bold_title="Google Chrome Extension (Manifest V3):")
    add_bullet_p(doc, " Plataforma PaaS de nuvem que provê sincronização de dados global em tempo real (WebSockets), suporte nativo a regras de segurança declarativas e autenticação robusta sem necessidade de gerenciar infraestruturas complexas de servidores de aplicação.", bold_title="Google Firebase (Auth & Cloud Firestore):")

    # =========================================================================
    # CAPÍTULO 6: MODELAGEM DE BANCO DE DADOS
    # =========================================================================
    print("[PARTE 3] Injetando Capítulo 6: Modelagem de Banco de Dados e Dicionário de Dados...")
    doc.add_page_break()
    
    add_heading_1(doc, "6. MODELAGEM DE BANCO DE DADOS")
    
    add_heading_2(doc, "6.1 Abordagem e Estrutura Conceitual (MER) e Lógica (DER)")
    add_body_p(
        doc,
        "A modelagem dos dados do SeguranApp foi projetada para combinar a integridade relacional exigida pela auditoria "
        "e governança corporativa com a flexibilidade de persistência NoSQL distribuída em documentos do Cloud Firestore. "
        "O Modelo Entidade-Relacionamento (MER) estrutura o ecossistema em torno da entidade central USUARIO, vinculada ao "
        "COFRE_SEGURANCA (relação 1:1), às CREDENCIAIS individuais (relação 1:N), às TRILHAS_AUDITORIA (relação 1:N), "
        "às SESSOES_ATIVAS (relação 1:N) e ao PLANO_ASSINATURA (relação 1:1)."
    )
    add_body_p(
        doc,
        "No DER Lógico, as entidades relacionam-se através de chaves primárias e estrangeiras consolidadas no documento do "
        "Firestore sob a raiz 'users/{userId}'. A chave 'userId' assegura o particionamento isolado dos dados de cada organização "
        "ou colaborador. As seções a seguir trazem o Dicionário de Dados exaustivo em tabelas padronizadas para cada uma das "
        "oito entidades do sistema."
    )
    
    add_heading_2(doc, "6.2 Dicionário de Dados Exaustivo do Sistema")
    
    # 6.2.1 USUARIOS_PERFIL
    add_heading_3(doc, "6.2.1 Entidade: USUARIOS_PERFIL (Perfil do Usuário e Metadados Criptográficos)")
    add_body_p(doc, "Armazena as credenciais de autenticação local e em nuvem, parâmetros de derivação KDF e configurações de segurança:")
    t_user_headers = ["Atributo", "Tipo de Dado", "Tamanho", "PK/FK", "Nulo", "Restrições e Descrição Semântica"]
    t_user_rows = [
        ["id", "VARCHAR", "128", "PK", "Não", "Identificador único universal (UID gerado pelo Firebase Authentication)."],
        ["username", "VARCHAR", "255", "-", "Não", "E-mail institucional ou nome de login único do usuário no sistema."],
        ["salt", "VARCHAR", "64", "-", "Não", "Sequência aleatória gerada por CSPRNG de 16 bytes codificada em Base64 para derivação."],
        ["kdfAlgorithm", "VARCHAR", "16", "-", "Não", "Algoritmo de derivação da Master Key ('pbkdf2' ou 'argon2id')."],
        ["passwordHash", "VARCHAR", "128", "-", "Não", "Hash SHA-256 da concatenação da senha mestra com o salt para verificação rápida."],
        ["pinHash", "VARCHAR", "128", "-", "Sim", "Hash SHA-256 do PIN local de 6 dígitos utilizado para desbloqueio rápido."],
        ["encryptedTotpSecret", "JSON / TEXT", "512", "-", "Sim", "Envelope criptográfico (ciphertext + iv) contendo o segredo do 2FA corporativo."],
        ["totpEnabled", "BOOLEAN", "1", "-", "Não", "Flag que indica se a autenticação em duas etapas por TOTP está ativa."],
        ["createdAt", "TIMESTAMP", "-", "-", "Não", "Carimbo de data e hora ISO 8601 do cadastramento do usuário na plataforma."],
    ]
    add_table_data(doc, t_user_headers, t_user_rows, [3.0, 2.0, 1.5, 1.3, 1.2, 7.0])

    # 6.2.2 COFRE_ENVELOPE
    add_heading_3(doc, "6.2.2 Entidade: COFRE_ENVELOPE (Envelope Criptográfico Raiz do Cofre)")
    add_body_p(doc, "Implementa o padrão Envelope Encryption, abrigando a Chave de Cifragem de Dados (DEK) cifrada:")
    t_env_headers = ["Atributo", "Tipo de Dado", "Tamanho", "PK/FK", "Nulo", "Restrições e Descrição Semântica"]
    t_env_rows = [
        ["id", "VARCHAR", "128", "PK", "Não", "Identificador do documento envelope do cofre no Cloud Firestore."],
        ["userId", "VARCHAR", "128", "FK", "Não", "Referência à chave primária da entidade USUARIOS_PERFIL."],
        ["version", "INTEGER", "4", "-", "Não", "Número da versão estrutural do cofre (utilizado para migrações de esquema)."],
        ["encryptedDekCiphertext", "TEXT", "512", "-", "Sim", "Chave DEK de 256 bits gerada aleatoriamente e cifrada com a KEK (AES-GCM 256)."],
        ["encryptedDekIv", "VARCHAR", "64", "-", "Sim", "Vetor de inicialização de 96 bits em Base64 utilizado na cifragem da DEK."],
        ["updatedAt", "TIMESTAMP", "-", "-", "Não", "Data/hora da última sincronização ou rotação do envelope criptográfico."],
    ]
    add_table_data(doc, t_env_headers, t_env_rows, [3.2, 2.0, 1.5, 1.3, 1.2, 6.8])

    # 6.2.3 COFRE_ITENS
    add_heading_3(doc, "6.2.3 Entidade: COFRE_ITENS (Itens de Credenciais Individuais)")
    add_body_p(doc, "Registra cada credencial, cartão corporativo, nota segura ou identidade cadastrada pelo colaborador:")
    t_item_headers = ["Atributo", "Tipo de Dado", "Tamanho", "PK/FK", "Nulo", "Restrições e Descrição Semântica"]
    t_item_rows = [
        ["id", "VARCHAR", "64", "PK", "Não", "Identificador UUID v4 exclusivo gerado localmente para o item do cofre."],
        ["envelopeId", "VARCHAR", "128", "FK", "Não", "Referência ao envelope proprietário da credencial (COFRE_ENVELOPE)."],
        ["title", "VARCHAR", "120", "-", "Não", "Nome amigável descritivo do recurso (ex: 'Portal Corporativo SAP', 'E-mail RH')."],
        ["username", "VARCHAR", "255", "-", "Não", "Identificador de acesso associado à credencial (e-mail, matrícula, login)."],
        ["encryptedPasswordCiphertext", "TEXT", "1024", "-", "Não", "Senha cifrada com a DEK via AES-GCM 256 bits codificada em Base64."],
        ["encryptedPasswordIv", "VARCHAR", "64", "-", "Não", "Vetor de inicialização de 96 bits exclusivo utilizado para cifrar a senha."],
        ["url", "VARCHAR", "500", "-", "Sim", "Endereço URL completo do sistema para cruzamento com o detector e autofill."],
        ["category", "VARCHAR", "20", "-", "Não", "Categoria restrita: 'LOGIN', 'CARD', 'NOTE', 'IDENTITY', 'OTHER'."],
        ["encryptedNotesCiphertext", "TEXT", "4096", "-", "Sim", "Anotações confidenciais e tokens de acesso secundários cifrados."],
        ["encryptedNotesIv", "VARCHAR", "64", "-", "Sim", "Vetor de inicialização específico utilizado no campo de anotações."],
        ["encryptedTotpSecret", "TEXT", "512", "-", "Sim", "Chave semente Base32 do autenticador de segundo fator individual cifrada."],
        ["updatedAt", "TIMESTAMP", "-", "-", "Não", "Carimbo de data e hora da última modificação efetuada nesta credencial."],
    ]
    add_table_data(doc, t_item_headers, t_item_rows, [3.2, 2.0, 1.3, 1.3, 1.2, 7.0])

    # 6.2.4 HISTORICO_SENHAS
    add_heading_3(doc, "6.2.4 Entidade: HISTORICO_SENHAS (Auditoria de Senhas Anteriores)")
    add_body_p(doc, "Preserva versões antigas de senhas para permitir auditoria e recuperação em caso de alterações indesejadas:")
    t_hist_headers = ["Atributo", "Tipo de Dado", "Tamanho", "PK/FK", "Nulo", "Restrições e Descrição Semântica"]
    t_hist_rows = [
        ["id", "VARCHAR", "64", "PK", "Não", "Identificador UUID do registro histórico da credencial."],
        ["itemId", "VARCHAR", "64", "FK", "Não", "Chave estrangeira apontando para o item correspondente em COFRE_ITENS."],
        ["username", "VARCHAR", "255", "-", "Não", "Nome de usuário que estava ativo durante a vigência desta senha histórica."],
        ["encryptedPasswordCiphertext", "TEXT", "1024", "-", "Não", "Senha anterior cifrada em AES-GCM 256 para manter o sigilo permanente."],
        ["encryptedPasswordIv", "VARCHAR", "64", "-", "Não", "Vetor de inicialização associado à cifragem da senha histórica."],
        ["updatedAt", "TIMESTAMP", "-", "-", "Não", "Data e hora exata em que a senha anterior foi substituída pela nova."],
    ]
    add_table_data(doc, t_hist_headers, t_hist_rows, [3.2, 2.0, 1.3, 1.3, 1.2, 7.0])

    # 6.2.5 LOGS_AUDITORIA
    add_heading_3(doc, "6.2.5 Entidade: LOGS_AUDITORIA (Trilha de Auditoria e Conformidade LGPD)")
    add_body_p(doc, "Registra eventos de segurança sensíveis e garante o princípio da responsabilização exigido pela LGPD:")
    t_log_headers = ["Atributo", "Tipo de Dado", "Tamanho", "PK/FK", "Nulo", "Restrições e Descrição Semântica"]
    t_log_rows = [
        ["id", "VARCHAR", "64", "PK", "Não", "Identificador UUID do evento de auditoria."],
        ["userId", "VARCHAR", "128", "FK", "Não", "Chave estrangeira que identifica o usuário que gerou a ação (USUARIOS_PERFIL)."],
        ["timestamp", "TIMESTAMP", "-", "-", "Não", "Carimbo de data e hora do evento no padrão ISO 8601 em fuso UTC."],
        ["action", "VARCHAR", "40", "-", "Não", "Código do evento (ex: 'LOGIN_SUCCESS', 'PIN_LOCKED', 'DEK_ROTATED', 'VPN_CONNECTED')."],
        ["details", "VARCHAR", "500", "-", "Não", "Descrição contextual do evento (ex: 'Conexão VPN estabelecida com sucesso no nó VPS SP')."],
    ]
    add_table_data(doc, t_log_headers, t_log_rows, [2.5, 2.0, 1.3, 1.3, 1.2, 7.7])

    # 6.2.6 SERVIDORES_VPN
    add_heading_3(doc, "6.2.6 Entidade: SERVIDORES_VPN (Nós de Conexão e Túneis WireGuard)")
    add_body_p(doc, "Armazena os parâmetros operacionais dos servidores VPN WireGuard homologados e instâncias customizadas:")
    t_vpn_headers = ["Atributo", "Tipo de Dado", "Tamanho", "PK/FK", "Nulo", "Restrições e Descrição Semântica"]
    t_vpn_rows = [
        ["id", "VARCHAR", "40", "PK", "Não", "Identificador do servidor VPN (ex: 'br-sp-01', 'cf-warp-01')."],
        ["name", "VARCHAR", "100", "-", "Não", "Nome amigável de exibição (ex: 'VPS São Paulo - Baixa Latência')."],
        ["country", "VARCHAR", "60", "-", "Não", "País de hospedagem da infraestrutura de rede."],
        ["city", "VARCHAR", "60", "-", "Não", "Cidade ou região geográfica onde a máquina virtual reside."],
        ["flag", "VARCHAR", "10", "-", "Não", "Caractere emoji correspondente à bandeira nacional para interface."],
        ["endpoint", "VARCHAR", "150", "-", "Não", "Endereço IPv4 estático ou FQDN público da máquina servidora WireGuard."],
        ["port", "INTEGER", "4", "-", "Não", "Porta UDP de escuta do serviço WireGuard (padrão oficial: 51820)."],
        ["serverPublicKey", "VARCHAR", "64", "-", "Não", "Chave pública Curve25519 (Base64 de 32 bytes) para troca de chaves Diffie-Hellman."],
        ["clientAddress", "VARCHAR", "32", "-", "Não", "Faixa de endereçamento IP privado atribuído à interface virtual do cliente (ex: 10.14.0.2/32)."],
        ["status", "VARCHAR", "20", "-", "Não", "Estado operacional do nó: 'AVAILABLE', 'UNAVAILABLE', 'MAINTENANCE'."],
        ["isCustom", "BOOLEAN", "1", "-", "Não", "Flag que diferencia servidores corporativos oficiais de servidores privados adicionados."],
        ["ping", "INTEGER", "4", "-", "Sim", "Latência média de ida e volta (RTT) calculada em milissegundos."],
    ]
    add_table_data(doc, t_vpn_headers, t_vpn_rows, [2.8, 1.8, 1.3, 1.3, 1.2, 7.6])

    # 6.2.7 ASSINATURAS_PLANOS
    add_heading_3(doc, "6.2.7 Entidade: ASSINATURAS_PLANOS (Gestão de Contratos e Níveis de Assinatura)")
    add_body_p(doc, "Controla os termos de licenciamento, ciclos de cobrança e permissões de acesso às funcionalidades do sistema:")
    t_sub_headers = ["Atributo", "Tipo de Dado", "Tamanho", "PK/FK", "Nulo", "Restrições e Descrição Semântica"]
    t_sub_rows = [
        ["id", "VARCHAR", "64", "PK", "Não", "Identificador do registro contratual de assinatura SaaS."],
        ["userId", "VARCHAR", "128", "FK", "Não", "Chave estrangeira apontando para USUARIOS_PERFIL."],
        ["tier", "VARCHAR", "16", "-", "Não", "Nível de serviço contratado: 'FREE', 'PRO' ou 'ULTRA' (Enterprise)."],
        ["billingCycle", "VARCHAR", "16", "-", "Sim", "Ciclo de faturamento: 'MONTHLY' ou 'ANNUAL'."],
        ["status", "VARCHAR", "16", "-", "Não", "Situação do contrato: 'ACTIVE', 'EXPIRED' ou 'TRIAL'."],
        ["startedAt", "TIMESTAMP", "-", "-", "Não", "Data e hora de ativação inicial do plano."],
        ["expiresAt", "TIMESTAMP", "-", "-", "Sim", "Data de expiração do acesso ou null para contas comunitárias vitalícias."],
        ["autoRenew", "BOOLEAN", "1", "-", "Não", "Flag que indica se a renovação automática da assinatura está habilitada."],
        ["paymentMethod", "VARCHAR", "20", "-", "Sim", "Método de liquidação: 'PIX' ou 'CREDIT_CARD'."],
    ]
    add_table_data(doc, t_sub_headers, t_sub_rows, [2.8, 2.0, 1.3, 1.3, 1.2, 7.4])

    # 6.2.8 SESSOES_ATIVAS
    add_heading_3(doc, "6.2.8 Entidade: SESSOES_ATIVAS (Controle de Dispositivos Conectados)")
    add_body_p(doc, "Monitora os aparelhos móveis e estações autenticadas para revogação remota em caso de furto ou extravio:")
    t_ses_headers = ["Atributo", "Tipo de Dado", "Tamanho", "PK/FK", "Nulo", "Restrições e Descrição Semântica"]
    t_ses_rows = [
        ["id", "VARCHAR", "64", "PK", "Não", "Identificador exclusivo do token de sessão ativo."],
        ["userId", "VARCHAR", "128", "FK", "Não", "Chave estrangeira apontando para USUARIOS_PERFIL."],
        ["deviceName", "VARCHAR", "100", "-", "Não", "Descrição do hardware (ex: 'Samsung Galaxy S23 Android 14', 'Chrome Windows 11')."],
        ["lastActive", "TIMESTAMP", "-", "-", "Não", "Carimbo de data/hora da última atividade registrada neste dispositivo."],
        ["ipPlaceholder", "VARCHAR", "45", "-", "Sim", "Indicação de IP protegido ou rótulo de túnel WireGuard ativo."],
    ]
    add_table_data(doc, t_ses_headers, t_ses_rows, [2.8, 2.0, 1.3, 1.3, 1.2, 7.4])

    # =========================================================================
    # CAPÍTULO 7: PROTOTIPAÇÃO E DESIGN DE INTERFACE (UI/UX E BEM-ESTAR)
    # =========================================================================
    print("[PARTE 3] Injetando Capítulo 7: Prototipação UI/UX e Guia de Telas...")
    doc.add_page_break()
    
    add_heading_1(doc, "7. PROTOTIPAÇÃO E DESIGN DE INTERFACE (UI/UX E BEM-ESTAR)")
    
    add_heading_2(doc, "7.1 Princípios de UI/UX Voltados à Redução de Erros Operacionais e Estresse Digital")
    add_body_p(
        doc,
        "A experiência do usuário (UX) em aplicações de segurança corporativa frequentemente peca pelo excesso de burocracia, "
        "gerando fricção, desmotivação e incentivo involuntário ao descumprimento de protocolos. No desenvolvimento do SeguranApp, "
        "a interface do usuário (UI) foi concebida sob os preceitos do Design Ergonômico e da Psicologia Comportamental, "
        "focada na redução da carga cognitiva dos colaboradores."
    )
    add_body_p(
        doc,
        "Para atingir esses objetivos, foram implementadas as seguintes diretrizes estéticas e funcionais:",
        bold_prefix="Diretrizes Ergonômicas do SeguranApp:\n"
    )
    add_bullet_p(doc, " A paleta de cores adota tons escuros suaves baseados na escala Slate (#0F172A, #1E293B) com contrastes calibrados segundo o padrão WCAG AAA. Essa escolha reduz significativamente a fadiga visual e o estresse ocular em ambientes de trabalho prolongados.", bold_title="Paleta Visual Calmante (Dark Mode Glassmorphism):")
    add_bullet_p(doc, " Utilização da família tipográfica moderna Inter e Roboto, reconhecidas pela excepcional legibilidade em telas de alta densidade de pixels (Retina e AMOLED), reduzindo o esforço de leitura de senhas alfanuméricas complexas.", bold_title="Tipografia Otimizada para Criptografia:")
    add_bullet_p(doc, " Eliminação de janelas modais agressivas e mensagens de alerta punitivas. O sistema emprega indicadores visuais graduais (verde para seguro, amarelo para atenção, vermelho para perigo) acompanhados de explicações construtivas e orientadoras que transformam incidentes em oportunidades de aprendizado.", bold_title="Prevenção da Fadiga de Alertas (Alert Fatigue):")
    add_bullet_p(doc, " Ações críticas (como cópia de senha mestra, geração de 2FA e ativação do túnel VPN) são executadas com no máximo dois toques na tela, reduzindo o tempo de exposição e a probabilidade de erros manuais.", bold_title="Eficiência em Dois Cliques:")

    add_heading_2(doc, "7.2 Protótipo Navegável Figma e Especificação Tela a Tela")
    add_body_p(
        doc,
        "O protótipo navegável em alta fidelidade foi desenvolvido no Figma, mapeando todos os estados de interação, "
        "transições animadas e fluxos responsivos para smartphone e extensão desktop. O acesso interativo ao protótipo "
        "pode ser realizado através do link oficial:",
        bold_prefix="Referência do Protótipo de Alta Fidelidade: "
    )
    add_body_p(
        doc,
        "https://www.figma.com/design/seguranapp-riskzero-tcc-senai (Protótipo Navegável Interativo)",
        bold_prefix="Link Figma: "
    )
    
    add_body_p(
        doc,
        "As seções a seguir trazem a especificação detalhada tela a tela do ecossistema, acompanhadas dos marcadores formais "
        "para inserção das capturas de tela finais:"
    )

    # TELAS COM PLACEHOLDERS
    screens_info = [
        (
            "Tela 1: Splash Screen & Onboarding Criptográfico 3D",
            "Apresenta a identidade visual da plataforma através de um modelo 3D dinâmico renderizado em tempo real com Spline 3D, "
            "representando o escudo criptográfico do SeguranApp. Apresenta mensagens de boas-vindas acolhedoras, orientações iniciais "
            "sobre a importância da custódia pessoal da Senha Mestra e navegação guiada para a criação ou restauração do cofre.",
            "[INSERIR CAPTURA DE TELA 1 - SPLASH SCREEN E ESCUDO CRIPTOGRÁFICO 3D]"
        ),
        (
            "Tela 2: Autenticação Mestra, Biometria e 2FA/TOTP",
            "Interface de liberação de acesso com foco em segurança sem atrito. Apresenta campo com máscara de proteção para a Senha Mestra, "
            "botão de ativação rápida por impressão digital/FaceID via Android BiometricPrompt e prompt de validação do código TOTP temporário "
            "de 6 dígitos. Exibe feedback imediato em caso de erro e bloqueio progressivo após tentativas inválidas.",
            "[INSERIR CAPTURA DE TELA 2 - TELA DE AUTENTICAÇÃO MESTRA, BIOMETRIA E TOTP]"
        ),
        (
            "Tela 3: Dashboard Principal (Home) com Risk Score e Status de Proteção",
            "Painel central que sintetiza a postura defensiva do colaborador através de um mostrador radial de Pontuação de Risco (Risk Score de 0 a 100). "
            "Exibe cards de acesso rápido ao Cofre, à VPN WireGuard, ao Detector de Links e à Central Educativa. Um carrossel informativo exibe pílulas "
            "rápidas de conscientização e dicas para manutenção do bem-estar e equilíbrio no uso da tecnologia corporativa.",
            "[INSERIR CAPTURA DE TELA 3 - DASHBOARD PRINCIPAL COM RISK SCORE E STATUS DE PROTEÇÃO]"
        ),
        (
            "Tela 4: Cofre de Senhas e Credenciais (Listagem e Detalhes Criptografados)",
            "Gerenciador central de credenciais com barra de pesquisa instantânea em memória decifrada e filtros por categoria (Login, Cartões, Notas). "
            "Cada registro exibe título, URL e usuário. Ao tocar na credencial, abre-se a visão detalhada com botão de revelação temporária da senha, "
            "cópia segura com contagem regressiva de 30 segundos, exibição do código 2FA correspondente e histórico de alterações.",
            "[INSERIR CAPTURA DE TELA 4 - COFRE DE CREDENCIAIS COM CÓPIA SEGURA E HISTÓRICO]"
        ),
        (
            "Tela 5: Conexão VPN WireGuard em Tempo Real (Status de IP e Túnel Ativo)",
            "Interface de controle do túnel criptográfico de rede. Apresenta seletor de servidores (VPS São Paulo, Cloudflare WARP ou nós customizados), "
            "grande botão tátil de alternância 'Conectar/Desconectar', medidores de latência (ping em ms), taxas de transferência de download/upload "
            "em tempo real e verificação em destaque comprovando o mascaramento do endereço IP público e proteção contra vazamento de DNS.",
            "[INSERIR CAPTURA DE TELA 5 - TELA DA VPN WIREGUARD COM MASCARAMENTO DE IP E MÉTRICAS]"
        ),
        (
            "Tela 6: Detector de Links Maliciosos e Análise Heurística Anti-Phishing",
            "Ferramenta defensiva com campo para colagem ou compartilhamento direto de URLs suspeitas recebidas por e-mail, WhatsApp ou SMS. "
            "Executa testes heurísticos em tempo real (HTTPS, @ spoofing, IP numérico, encurtador, TLD malicioso) e apresenta um relatório visual detalhado "
            "com Threat Score (0 a 100), veredito semafórico (Seguro, Suspeito, Perigoso) e histórico das últimas auditorias realizadas.",
            "[INSERIR CAPTURA DE TELA 6 - DETECTOR HEURÍSTICO DE LINKS SUSPEITOS E THREAT SCORE]"
        ),
        (
            "Tela 7: Auditoria de Vazamentos de Dados (k-Anonymity HIBP)",
            "Módulo de verificação proativa de exposição de identidade. Permite ao colaborador consultar se seu e-mail corporativo constou em violações "
            "de bases de dados conhecidas e testar a robustez de senhas através do modelo matemático k-Anonymity, garantindo que o texto plano jamais "
            "deixe o celular. Exibe relatórios detalhados com datas, contexto do incidente e orientações claras de remediação.",
            "[INSERIR CAPTURA DE TELA 7 - AUDITORIA DE VAZAMENTOS VIA K-ANONYMITY]"
        ),
        (
            "Tela 8: Módulo de Hardening do Dispositivo e Exportação de Logs LGPD",
            "Painel de segurança nativa do aparelho. Permite ativar a proteção Privacy Screen (bloqueio de prints com FLAG_SECURE), demonstrar a limpeza "
            "do clipboard em 30 segundos, ativar o modo de camuflagem de interface e visualizar a trilha imutável de logs de auditoria, "
            "com botão de exportação instantânea de relatório de conformidade em formato CSV compatível com os requisitos da LGPD.",
            "[INSERIR CAPTURA DE TELA 8 - PAINEL DE HARDENING DO DISPOSITIVO E EXPORTAÇÃO CSV LGPD]"
        ),
        (
            "Tela 9: Central de Aprendizado em Segurança e Quiz de Conscientização",
            "Ambiente de microaprendizagem corporativa estruturado em abas dinâmicas: Ataques Comuns, Como Evitar e Boas Práticas (totalizando 13 tópicos ricos). "
            "Inclui o Quiz de Privacidade e Segurança com perguntas conceituais, barra de progresso interativa, feedback imediato e medalhas de aproveitamento, "
            "promovendo o aprendizado contínuo sem estresse avaliativo.",
            "[INSERIR CAPTURA DE TELA 9 - CENTRAL DE EDUCAÇÃO DIGITAL E QUIZ DE PRIVACIDADE]"
        ),
        (
            "Tela 10: Extensão para Navegadores Chrome (Autofill Seguro e Detecção de Abas)",
            "Interface compacta da extensão Google Chrome Manifest V3. Detecta automaticamente o domínio da aba ativa, preenche credenciais em formulários "
            "com um clique, avalia a reputação da URL visitada em tempo real e oferece controle direto para ativar a proteção de rede VPN.",
            "[INSERIR CAPTURA DE TELA 10 - EXTENSÃO CHROME MANIFEST V3 COM AUTOFILL E VERIFICAÇÃO]"
        ),
    ]

    for title, desc, placeholder in screens_info:
        add_heading_3(doc, title)
        add_body_p(doc, desc)
        
        # Inserção do placeholder com destaque visual
        p_ph = doc.add_paragraph()
        p_ph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_ph.paragraph_format.space_before = Pt(8)
        p_ph.paragraph_format.space_after = Pt(12)
        r_ph = p_ph.add_run(placeholder)
        r_ph.bold = True
        r_ph.font.name = "Arial"
        r_ph.font.size = Pt(10)
        r_ph.font.color.rgb = RGBColor(0x25, 0x63, 0xEB) # Azul Royal Destacado

    # Salvar o documento atualizado
    doc.save(target_path)
    try:
        doc.save(workspace_target)
    except Exception as e:
        print(f"Aviso ao salvar na raiz do workspace: {e}")
        
    print(f"[PARTE 3 CONCLUÍDA COM SUCESSO!]")
    print(f"Arquivo gerado: {target_path}")
    print(f"Total de parágrafos no documento após Parte 3: {len(doc.paragraphs)}")
    print(f"Total de tabelas no documento após Parte 3: {len(doc.tables)}")

if __name__ == "__main__":
    run_part3()
